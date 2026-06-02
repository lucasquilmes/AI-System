"""
Genera 3 documentos Word para el estudio con usuarios finales.
Diseño contrabalanceado: cada persona ve cada texto en UNA sola versión (original o adaptada).

Asignación:
  Persona 1 (S1): T0-O, T1-A, T2-O, T3-A, T4-O  |  (S2): T5-A, T6-O, T7-A, T8-O, T9-A
  Persona 2 (S1): T1-O, T2-A, T3-O, T4-A, T5-O  |  (S2): T6-A, T7-O, T8-A, T9-O, T10-A
  Persona 3 (S1): T0-A, T2-O, T4-A, T6-O, T8-A  |  (S2): T10-O, T1-A, T3-O, T5-A, T7-O
"""

import pandas as pd
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

ROOT = Path(__file__).parent

# ── Cargar datos ───────────────────────────────────────────────────────────────

df = pd.read_excel(ROOT / "data/exemples_lectura_facil_formatted.xlsx", sheet_name="Hoja1")

jsonl = ROOT / "outputs/fase2/exemples_lectura_facil_formatted_t0.0_s42/v8/llama3.3/results.jsonl"
adapted = {}
for line in jsonl.read_text(encoding="utf-8").strip().split("\n"):
    r = json.loads(line)
    if "error" not in r:
        result = r["result"]
        text = result.get("raw_output", str(result)) if isinstance(result, dict) else str(result)
        adapted[int(r["item_id"])] = text.strip()

texts = []
for i, row in df.iterrows():
    texts.append({
        "idx":    i,
        "title":  str(row.get("title", "")).strip(),
        "orig":   str(row.get("normal text", "")).strip(),
        "adapt":  adapted.get(i, ""),
    })

# ── Diseño contrabalanceado ────────────────────────────────────────────────────
# Tuplas (idx_texto, "orig"/"adapt")

SESSIONS = {
    1: {
        1: [(0,"orig"),(1,"adapt"),(2,"orig"),(3,"adapt"),(4,"orig")],
        2: [(5,"adapt"),(6,"orig"),(7,"adapt"),(8,"orig"),(9,"adapt")],
    },
    2: {
        1: [(1,"orig"),(2,"adapt"),(3,"orig"),(4,"adapt"),(5,"orig")],
        2: [(6,"adapt"),(7,"orig"),(8,"adapt"),(9,"orig"),(10,"adapt")],
    },
    3: {
        1: [(0,"adapt"),(2,"orig"),(4,"adapt"),(6,"orig"),(8,"adapt")],
        2: [(10,"orig"),(1,"adapt"),(3,"orig"),(5,"adapt"),(7,"orig")],
    },
}

QUESTIONS = [
    "El text és fàcil de llegir.",
    "He entès bé el text.",
    "Hi havia paraules que no coneixia.",
    "Hi havia frases llargues o difícils d'entendre.",
]

# ── Helpers de formato ─────────────────────────────────────────────────────────

def set_font(run, size=12, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    if level == 1:
        set_font(run, size=20, bold=True, color=(0, 70, 127))
    elif level == 2:
        set_font(run, size=15, bold=True, color=(0, 70, 127))
    else:
        set_font(run, size=13, bold=True)
    return p


def add_body(doc, text, size=12, italic=False, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size=size, color=color)
    run.font.italic = italic
    return p


def add_question_block(doc, questions):
    """Añade las 4 preguntas con columnas SÍ / NO."""
    doc.add_paragraph()
    table = doc.add_table(rows=len(questions) + 1, cols=3)
    table.style = "Table Grid"

    # Cabecera
    hdr = table.rows[0].cells
    hdr[0].text = "Pregunta"
    hdr[1].text = "Sí"
    hdr[2].text = "No"
    for cell in hdr:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(11)
        cell._tc.get_or_add_tcPr()

    # Filas de preguntas
    for i, q in enumerate(questions):
        row = table.rows[i + 1].cells
        row[0].text = q
        row[1].text = "☐"
        row[2].text = "☐"
        for cell in row:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
        # Centrar las columnas Sí/No
        for col in [1, 2]:
            for p in row[col].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ancho de columnas
    for row in table.rows:
        row.cells[0].width = Cm(11)
        row.cells[1].width = Cm(2)
        row.cells[2].width = Cm(2)

    doc.add_paragraph()


def add_cover(doc, person_num):
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ESTUDI DE LECTURA FÀCIL")
    set_font(run, size=26, bold=True, color=(0, 70, 127))

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Participant {person_num}")
    set_font(run, size=20, bold=True)

    doc.add_paragraph()
    doc.add_paragraph()

    lines = [
        "Gràcies per participar en aquest estudi.",
        "",
        "Llegiràs una sèrie de textos.",
        "Després de cada text,",
        "respondràs 4 preguntes senzilles.",
        "",
        "Respondràs SÍ o NO a cada pregunta.",
        "",
        "No hi ha respostes correctes ni incorrectes.",
        "Ens interessa la teva opinió.",
        "",
        "Si tens algun dubte,",
        "demana ajuda a la persona que t'acompanya.",
    ]
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        set_font(run, size=13)

    doc.add_page_break()


def add_session(doc, person_num, session_num, session_items):
    # Encabezado de sesión
    add_heading(doc, f"SESSIÓ {session_num}", level=1)
    add_body(doc, f"Participant {person_num}  ·  Sessió {session_num} de 2", size=11, color=(100, 100, 100))
    doc.add_paragraph()

    for pos, (tidx, version) in enumerate(session_items, 1):
        t = texts[tidx]
        title_clean = t["title"].split("(")[0].strip()  # eliminar descripción entre paréntesis
        # Quitar número inicial si lo hay
        if title_clean and title_clean[0].isdigit():
            title_clean = ". ".join(title_clean.split(". ")[1:]).strip()

        # Separador de texto
        add_heading(doc, f"Text {pos} de 5", level=2)
        if title_clean:
            add_body(doc, title_clean, size=11, italic=True, color=(80, 80, 80))

        doc.add_paragraph()

        # Contenido del texto
        content = t["orig"] if version == "orig" else t["adapt"]
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.right_indent = Cm(0.5)
        run = p.add_run(content)
        set_font(run, size=12)

        doc.add_paragraph()

        # Preguntas
        add_body(doc, "Ara respon les preguntes:", size=12)
        add_question_block(doc, QUESTIONS)

        if pos < len(session_items):
            doc.add_paragraph()
            p = doc.add_paragraph("· · · · · · · · · · · · · · · · · · · · · · · · ·")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()

    doc.add_page_break()


# ── Generar documentos ─────────────────────────────────────────────────────────

out_dir = ROOT / "user_study"
out_dir.mkdir(exist_ok=True)

for person in [1, 2, 3]:
    doc = Document()

    # Márgenes
    section = doc.sections[0]
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(3)
    section.right_margin  = Cm(3)

    # Fuente por defecto
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(12)

    add_cover(doc, person)

    for session in [1, 2]:
        add_session(doc, person, session, SESSIONS[person][session])

    path = out_dir / f"participant_{person}.docx"
    doc.save(path)
    print(f"Guardado: {path}")

print("\nDiseño de asignación:")
for person in [1, 2, 3]:
    print(f"\n  Persona {person}:")
    for session in [1, 2]:
        items = SESSIONS[person][session]
        desc = ", ".join(f"T{idx}({'O' if v=='orig' else 'A'})" for idx, v in items)
        print(f"    Sesión {session}: {desc}")
